# -*- coding: utf-8 -*-
"""C 轮 DOCX：与 PDF 同源内容，插入真实图片与 C 轮截图。"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
D = os.path.join(BASE, 'delivery')
OUT = os.path.join(D, '02_明白办_产品说明书_可编辑源文件.docx')
GREEN = RGBColor(0x25, 0x4C, 0x40)
INK = RGBColor(0x26, 0x31, 0x2C)
GREY = RGBColor(0x5D, 0x6A, 0x63)
VER = '1.3.0-c7'

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.left_margin = s.right_margin = Cm(2.0)
s.top_margin = Cm(2.0); s.bottom_margin = Cm(1.8)
st = doc.styles['Normal']
st.font.name = '微软雅黑'; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
st.paragraph_format.line_spacing = 1.5
st.paragraph_format.space_after = Pt(4)


def cn(r, name='微软雅黑'):
    r.font.name = name
    r._element.rPr.rFonts.set(qn('w:eastAsia'), name)


def h(text, size=15, color=INK, before=12, after=6):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = color; cn(r)


def para(text, size=10.5, color=INK, bold=False):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text); r.font.size = Pt(size); r.font.color.rgb = color; r.bold = bold; cn(r)


def bullet(text, size=10):
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); r.font.size = Pt(size); cn(r)


def table(headers, rows, widths=None, font=9):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'
    for i, hd in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ''
        r = c.paragraphs[0].add_run(hd); r.bold = True; r.font.size = Pt(font); cn(r)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(font); cn(r)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def img(name, width_cm=16, cap=None):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(D, name), width=Cm(width_cm))
    if cap:
        c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(cap); r.font.size = Pt(8.5); r.font.color.rgb = GREY; cn(r)


# ---------- 封面 ----------
para('2026 iCAN 大学生创新创业大赛 · 快手 AI 未来创造者挑战赛　|　AI 产品实现', 9, GREY)
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10)
r = p.add_run('明白办'); r.bold = True; r.font.size = Pt(30); r.font.color.rgb = GREEN; cn(r)
para('给家人一张明白卡，也能读给他听', 14, INK, True)
para('把社区通知——文字或截图——整理成家人看得懂、能听清的行动卡：时间分用途、地点不混淆、每一项都能对照原文。', 11)
img('assets/illustrations/checking-together.png', 15.5, '情境插画为 AI 生成，人物为虚构角色，不代表真实个人。')
table(['项目', '内容'], [
    ['产品名称', '明白办（Mingbai Ban）'],
    ['版本', VER + '（三轮升级最终候选）'],
    ['在线体验', 'mingbaiban-42482.app.workbuddy.host'],
    ['运行时模型', 'deepseek-v4.1-flash（平台托管通道；取字与整理两次独立调用）'],
    ['本版新增', '截图取字与人工核对、行动卡朗读'],
], widths=[3.2, 13.0])
para('本说明书所述功能均为该版本已实现并可复现的功能；用户的真实使用效果尚未验证，第 7 节如实说明。', 9, GREY)
doc.add_page_break()

# ---------- 1 产品概述 ----------
h('01　产品概述：一件具体的小事', 15, GREEN)
para('下面这条通知是很多社区最常见的写法（本项目编写的演示样例，不是真实社区公告）：', 10.5)
table(['通知原文'], [['各位居民：为帮助老年居民掌握智能手机常用功能，社区服务中心将开展“手机使用指导”活动。时间：2026年10月26日9:00—11:00　地点：社区服务中心二楼活动室　报名方式：请于2026年10月25日17:00前到社区服务中心一楼服务台报名，名额有限。说明：本次活动仅限已报名居民参加，请携带本人手机。']], widths=[16])
table(['容易读错的地方', '为什么'], [
    ['两个时间', '25 日 17:00 是报名截止，26 日 9:00 是活动开始。错当成“25 日去参加”，就白跑一趟。'],
    ['两处地点', '报名在一楼服务台，活动在二楼活动室。合并成一个“社区服务中心”，到了还得再问。'],
    ['一个限制', '“仅限已报名居民参加”决定了要不要提前跑一趟；漏掉这句，当天可能进不去。'],
    ['一份材料', '只写了“本人手机”。实际准备时，人容易顺手再带身份证、医保卡——通知并没有要求。'],
], widths=[4.2, 11.8])
para('明白办做的事：把这条通知整理成一张卡——先说这条通知要办什么，再把两个时间、两处地点、限制条件、要带的东西分条列出，每一条都能点开看到通知里的原句；通知没写的，就写“通知未说明”，不替它补。C 轮之后，通知是截图也能直接传：AI 先取字、家人对照原图核对，卡片还能用设备语音读给长辈听。', 10.5)
para('目标用户：主要操作人是家人或社区志愿者，受益者是长辈。长辈不必自己复制粘贴长段文字；但卡片写得中性、可独立阅读，去不去、怎么办仍由本人决定。', 10)
doc.add_page_break()

# ---------- 2 用户与流程 ----------
h('02　目标用户与协作流程', 15, GREEN)
img('11_公众版五步图.png', 16, '图 1　一条通知到一张能听的卡（可编辑源：sources/flow-public.html）')
table(['角色', '具体动作', '为什么必须这样分'], [
    ['家人 / 志愿者', '贴文字或传截图、对照原图核对取出的文字、勾选“我已核对”、把卡读给长辈或转发', '只有人知道这条通知和自家情况的关系；核对是人的操作，不是社区认证'],
    ['AI · 取字', '只把截图里的文字逐字取出来，不改写、不补全；看不清标注', '整理者不用照着图打一遍字'],
    ['AI · 整理', '读懂任意写法、拆出事实、判断缺失与冲突，输出固定 JSON', '通知措辞千变万化，规则与模板覆盖不了'],
    ['程序', '结构、字段、每条引文是否出现在“核对后的文字”里', '机械检查每次都一样，不依赖模型心情'],
    ['长辈', '听卡或看卡、决定去不去、按卡准备', '——'],
], widths=[3.0, 7.6, 5.4])
h('适老化不是“把字放大”', 11, before=6)
bullet('正文 20px、行高 1.75；关键事实 22px；按钮最小 56px；复选框 32px、整行可点。')
bullet('状态不只靠颜色：原文写明 ✓ 实线、通知未说明 — 虚线、说得含糊 ? 点线、前后不一致 ! 双线。')
bullet('360/390/768/1280 实测无横向滚动；核对页手机上下排、桌面原图与文字并排。')
bullet('朗读由点击触发，不自动播放；可暂停、继续、停止；同一时刻只有一个朗读任务。')
doc.add_page_break()

# ---------- 3 核心功能 ----------
h('03　核心功能：六项，都在真实界面里', 15, GREEN)
table(['编号', '功能', '实现要点'], [
    ['F1', '通知输入', '粘贴文字（≤5000 字）或上传一张截图（JPG/PNG/WebP，≤8MB，长边缩到 1600px）；坏文件、超长都明确提示且不静默截断'],
    ['F2', '截图取字', '多模态模型独立调用，只取字不改写；看不清标注“［此处看不清，请对照图片］”；无字、太模糊如实提示'],
    ['F3', '人工核对', '左边原图、右边文字逐行对照；必须勾选“我已核对”才能整理；改字自动取消勾选，旧卡/复制/朗读全部作废'],
    ['F4', '事实整理', '一句话说清“要办什么”，再按 谁适用 / 何时何地 / 准备什么 / 下一步 / 还要确认 分条；冲突置顶'],
    ['F5', '原文核对', '文字来源标“对照通知原文”，截图来源标“对照核对后的通知文字”并可查看原图'],
    ['F6', '朗读与复制', '设备中文语音整卡朗读（暂停/继续/停止）；复制文本与勾选一致并声明“不是报名凭证”'],
], widths=[1.3, 2.7, 12])
img('assets/screenshots/c-02-upload-fig.png', 5.6, '① 传截图：先预览再识别')
img('assets/screenshots/c-03-review-fig.png', 5.6, '② 原图与文字对照，勾“我已核对”')
img('assets/screenshots/c-07-speech-fig.png', 5.6, '③ 读给家人听：设备语音，可暂停')
para('三处关键设计：①核对是闸门——不勾选就不能整理，证据入口写“核对后的文字”；②冲突或含糊置顶提醒；③朗读文本由已校验数据确定性拼出，不再让模型改写，也不会把勾选读成“已报名”。', 9.5, GREY)
doc.add_page_break()

# ---------- 4 场景示例 ----------
h('04　应用场景示例：原文与行动卡逐项对照', 15, GREEN)
table(['通知原文', '行动卡怎么显示', '关键点'], [
    ['请于 2026 年 10 月 25 日 17:00 前到社区服务中心一楼服务台报名', '报名截止：2026 年 10 月 25 日 17:00 前（原文写明）；地点：一楼服务台（原文写明）', '时间与它自己的地点绑在一起'],
    ['时间：2026 年 10 月 26 日 9:00—11:00　地点：二楼活动室', '活动开始：2026 年 10 月 26 日 9:00—11:00（原文写明）；地点：二楼活动室（原文写明）', '不与报名地点混成一个'],
    ['本次活动仅限已报名居民参加', '谁适用：已报名的居民（原文写明）', '限制条件不被删掉'],
    ['请携带本人手机', '准备什么：本人手机（原文写明）', '不加身份证、不加医保卡'],
    ['（原文没有写集合点、联系方式等）', '显示“通知未说明”，并在“还要确认什么”列出可问的问题', '缺就是缺，不补齐'],
], widths=[5.6, 7.4, 3])
img('assets/screenshots/c-08-missing-fig.png', 5.4, '缺信息：标“通知未说明”')
img('assets/screenshots/c-09-conflict-fig.png', 5.4, '冲突：顶部先提醒确认')
img('assets/screenshots/c-06-evidence-fig.png', 5.4, '对照核对后的通知文字')
para('以上均为线上版本真实返回与真实界面截图，未做内容修改。示例通知为本项目编写的虚构演示材料。', 9, GREY)
doc.add_page_break()

# ---------- 5 AI 的角色与价值 ----------
h('05　AI 的角色与价值', 15, GREEN)
table(['环节', '谁做', '边界'], [
    ['截图取字', '多模态模型（deepseek-v4.1-flash，独立一次调用、专用取字提示词）', '不改写、不补全、不给置信度；图里的指令只当文字'],
    ['人工核对', '整理者：对照原图改字、勾选“我已核对”', '是整理者的操作，不等于社区认证'],
    ['事实整理', '模型（同通道第二次调用）：固定 JSON 输出', '必须给逐字原文摘录'],
    ['机械校验', '程序：字段、状态、引文回溯、required 必须有依据', '校验不过→重试 1 次→仍不过不出卡'],
    ['朗读', '程序拼文本 + 设备语音发声（Web Speech API）', '不是云端 TTS；音色取决于设备'],
    ['判断去不去', '人：长辈听完/看完自己决定', '产品不代办、不报名'],
], widths=[3.0, 8.0, 5.0])
h('为什么朗读不让模型再写一遍', 11, before=6)
para('朗读文本从已通过校验的同一份卡片数据拼出来：数字、时点、地点与页面完全一致；“通知未说明”“前后不一致”“还要确认”都照读。时间只做等价改写（9:00 读作“9点”），数值、年份、用途不变；没有年份就不读年份。', 10)
para('必须说清的限制：引文能在核对后的文字里找到，不等于理解一定正确；取字也可能有错漏——所以核对必须由人完成。结果页始终提示以通知原文和社区说法为准。', 9.5, GREY)
table(['项目', '实际情况'], [
    ['接入方式', 'WorkBuddy 云服务免密钥通道，浏览器直连，前端不保存长期密钥'],
    ['实际模型', 'deepseek-v4.1-flash（取字与整理同模型、两次调用）'],
    ['朗读', '浏览器 Web Speech API + 设备中文音色（本机 Microsoft Huihui）；非云端 TTS'],
    ['无法确认', '上游机房形态与平台侧数据处理，本项目无法核实，故不作声明'],
], widths=[3.4, 12.6])
doc.add_page_break()

# ---------- 6 技术实现 ----------
h('06　技术实现思路', 15, GREEN)
img('11_技术实现图.png', 16, '图 2　技术实现与数据流（可编辑源：sources/flow-tech.html）')
h('技术选型', 11, before=6)
bullet('前端：原生 HTML/CSS/JS，无框架、无构建')
bullet('模型接入：WorkBuddy 云服务 SDK（CDN 版）')
bullet('取字：canvas 缩放 + base64 → OpenAI 兼容 image_url 消息（实测后接入）')
bullet('校验：自研 validate.js（结构与引文回溯）；朗读：Web Speech API')
h('关键实现', 11, before=6)
bullet('两套提示词各存一份：prompts/ocr-prompt.txt、prompts/system-prompt.txt')
bullet('请求序号丢弃过期响应；取字中可取消，换图不串台')
bullet('改字自动取消“已核对”，旧卡/复制/朗读一并作废')
bullet('全部文本走 textContent，不拼 innerHTML；图片只存于页面内存')
h('输出契约（节选）', 11, before=6)
para('document_status · summary · audience · times[{purpose,value,status,evidence,place}] · location · preparation · actions[{type,value,evidence}] · questions · warnings；status 仅取 stated/missing/unclear/conflict；missing 时 value 为 null、evidence 为空。', 9.5, GREY)
doc.add_page_break()

# ---------- 7 验证与限制 ----------
h('07　验证情况与限制', 15, GREEN)
table(['项目', '结果', '说明'], [
    ['C1 截图识别专项 20 项', '20 / 20 通过', '清晰/无年份无地点/冲突/模糊/无字/非通知/人工改字/图内指令/超长/个人信息不落库/坏文件/取消换图/核对闸门/超限不截断'],
    ['C2 朗读专项 19 项', '19 / 19 通过', '发音等价改写、缺失冲突照读、勾选不读成已报名、控制状态机、单实例、无音色与不支持的如实提示'],
    ['朗读原声核对', '3 张卡已实录', 'WASAPI 环回录下设备真实播放（完整/缺失/冲突），逐段测时长与音量'],
    ['A 轮回归（12 条用例）', '见 09 测试报告', '本版重跑，走真实模型'],
    ['页面行为 10 项 / A 轮专项 11 项', '见 09 测试报告', '本版重跑'],
    ['布局与可访问性', '通过', '360/390/768/1280 三个页面状态无横向滚动；触控 ≥44px；配色最低 5.12:1'],
], widths=[4.6, 3.6, 7.8])
h('还没做到的', 11, before=6)
bullet('未开展目标用户试用：没有效果、满意度、节省时间的数据。')
bullet('未做真机测试：微信内置浏览器、iOS Safari 的朗读兼容未验证。')
bullet('朗读依赖设备：音色随设备变化；不能导出音频文件，产品里没有“下载语音”按钮。')
bullet('取字不是认证：核对必须由人完成；校验拦不住“引文真实但理解错”。')
bullet('相对时间不换算；依赖网络与服务配额，被限流时如实报错。')
doc.add_page_break()

# ---------- 8 价值与下一步 ----------
h('08　应用价值与下一步', 15, GREEN)
table(['场景', '产品做的事'], [
    ['家人替长辈读通知', '截图直接传，不用照着打字；核对后转发，或当场读给长辈听'],
    ['长辈看字吃力', '整卡朗读，数字放慢读清；“未说明”“冲突”也照读'],
    ['社区志愿者协助登记', '写法各异的通知快速得到统一格式的事实清单，缺口与冲突一眼看出'],
    ['长辈自己复查', '卡片可独立阅读、可反复听；去不去仍由本人决定'],
], widths=[4.6, 11.4])
para('下一步怎么验证（不承诺结果）：配对试用（3—5 组“家人+长辈”，一条直接看、一条用产品，交叉顺序；记录答对率、是否听清“未说明/待确认”）；真机测试；若平台未来提供云端 TTS，再评估“保存这张卡的语音”。模板见 15_用户试用材料。', 10)
para('商业化与合作：目前没有任何商业合作、政府采购或社区落地；可能路径只是待验证设想。', 10)
doc.add_page_break()

# ---------- 9 来源与入口 ----------
h('09　用了什么工具，怎么自己试', 15, GREEN)
table(['用途', '工具 / 通道'], [
    ['产品运行时（取字）', 'deepseek-v4.1-flash 原生多模态（WorkBuddy 云服务通道）'],
    ['产品运行时（整理）', '同通道第二次独立调用，system-prompt.txt'],
    ['产品内朗读', '浏览器 Web Speech API + 设备中文音色；非云端 TTS，未克隆真人声音'],
    ['视频旁白', 'Windows SAPI5（Microsoft Huihui）；与产品内朗读是两个通道，分开标注'],
    ['朗读原声录制', 'WASAPI 环回录音（pyaudiowpatch）+ 有界面 Edge，录系统真实播放声'],
    ['插画', 'WorkBuddy 图像生成工具（沿用 B 轮 5 张，人物虚构）'],
    ['界面截图/录屏', 'Microsoft Edge + playwright-core（CDP 屏幕捕获）'],
], widths=[4.6, 11.4])
para('在线体验：mingbaiban-42482.app.workbuddy.host（无需注册）。控制台执行 window.MBB_CONFIG.appVersion 应返回 ' + VER + '。', 10)
para('声明：文中示例通知均为本项目编写的虚构演示材料；插画人物为虚构角色；本作品未使用第三方图库、音乐或他人成品；未声称技术首创，也未承诺晋级。', 9, GREY)

doc.save(OUT)
print('DOCX saved:', OUT, round(os.path.getsize(OUT) / 1024, 1), 'KB')
