# -*- coding: utf-8 -*-
"""生成 02_明白办_产品说明书_可编辑源文件.docx（内容与 PDF 说明书一致）"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(BASE, "delivery", "02_明白办_产品说明书_可编辑源文件.docx")
CHART = os.path.join(BASE, "delivery", "11_明白办_技术流程图.png")

BRAND = RGBColor(0x14, 0x53, 0x9A)
INK = RGBColor(0x1B, 0x1F, 0x24)
GREY = RGBColor(0x6B, 0x74, 0x80)

doc = Document()

# 页面与默认字体
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.left_margin = Cm(2.2)
sec.right_margin = Cm(2.2)
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.0)

style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.space_after = Pt(4)


def set_cn(run, name="微软雅黑"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def h(text, level=1, color=INK, size=None, space_before=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.bold = True
    r.font.color.rgb = color
    r.font.size = Pt(size or (16 if level == 1 else 12))
    set_cn(r)
    return p


def para(text, size=10.5, color=INK, bold=False, indent=0, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.first_line_indent = Cm(indent)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.bold = bold
    set_cn(r)
    return p


def bullet(text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.font.size = Pt(size)
    set_cn(r)
    return p


def table(headers, rows, widths=None, font=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        r = p.add_run(htext)
        r.bold = True
        r.font.size = Pt(font)
        set_cn(r)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            r = p.add_run(str(val))
            r.font.size = Pt(font)
            set_cn(r)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ================= 封面 =================
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(60)
r = p.add_run("2026 年 iCAN 大学生创新创业大赛 · 快手 AI 未来创造者挑战赛")
r.font.size = Pt(10)
r.font.color.rgb = GREY
set_cn(r)
p = doc.add_paragraph()
r = p.add_run("方向：AI 产品实现　|　主题：AI 连接的世界—公共服务与社会关怀")
r.font.size = Pt(10)
r.font.color.rgb = GREY
set_cn(r)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(36)
r = p.add_run("明白办")
r.bold = True
r.font.size = Pt(40)
r.font.color.rgb = BRAND
set_cn(r)

p = doc.add_paragraph()
r = p.add_run("面向长辈的 AI 社区通知行动助手")
r.font.size = Pt(15)
r.font.color.rgb = INK
set_cn(r)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
r = p.add_run("收到通知之后，下一步更明白。")
r.font.size = Pt(12)
r.font.color.rgb = BRAND
r.bold = True
set_cn(r)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(30)
r = p.add_run("产品说明书　v1.0.0（候选提交版）")
r.font.size = Pt(13)
r.bold = True
set_cn(r)

para("把一条看不懂的社区通知，变成一张能照着做的行动卡：说清是什么、谁能参加、什么时候、要带什么、下一步做什么，并且每个事实都能点开看原文依据。", size=11, space_after=10)

table(["项目", "内容"], [
    ["产品名称", "明白办（Mingbai Ban）"],
    ["版本", "v1.0.0（候选提交版）"],
    ["在线体验", "https://mingbaiban-42482.app.workbuddy.host/"],
    ["运行时模型", "WorkBuddy 云服务托管通道，实际调用 DeepSeek-V4.1-Flash（模型列表实测）"],
    ["文档日期", "2026 年 10 月 23 日"],
    ["文档性质", "参赛产品说明书；已验证结论与尚未验证的假设分别标注"],
], widths=[3.2, 13.0])

para("说明：文中所述功能均为当前线上版本已实现且可复现的功能。「已做测试」为真实执行记录；「目标用户效果」尚未开展真人验证，已在文末如实说明。文中截图均取自线上真实版本的实际运行画面。", size=9, color=GREY)
doc.add_page_break()

# ================= 一、产品概述 =================
h("一、产品概述", 1, BRAND)
para("明白办是一个网页工具。长辈把社区通知原文粘贴进去，它会用大白话输出一张行动卡：这件事是什么、谁能参加、分用途的时间、地点、要带什么、下一步做什么、还有哪些地方通知其实没说清。", size=11)

h("1.1 要解决的问题", 2)
para("社区通知写给「能读懂的成年人」，但对不少长辈，一条通知里至少藏着四道坎：")
table(["困难", "具体表现", "明白办的应对"], [
    ["时间说不清", "「25 日前报名」「28 日开始」混在一段里，容易把报名截止当成活动时间", "按用途拆开：报名截止 / 活动开始 / 服务时间，各自单列"],
    ["不知道要带什么", "「请携带本人手机」后面没有下文，长辈却会习惯性带上身份证", "只列原文写明的准备事项；没写就明确说「通知没提要带什么」"],
    ["不敢信、也没法核", "AI 一改写，就分不清哪句是原文、哪句是理解", "每个事实项可展开「看原文依据」，引文逐字回溯原文"],
    ["看完还是不动", "读懂了，但不知道先做什么、要不要问人", "可勾选的准备清单 + 待确认问题 + 一键复制给家人"],
], widths=[3.0, 7.0, 6.2])

h("1.2 主流程", 2)
para("① 粘贴文字通知 → ② 真实大模型理解 → ③ 白话行动卡（可展开原文依据）→ ④ 勾选准备事项 / 复制给家人。")

h("1.3 与常见「AI 摘要」的区别", 2)
para("普通摘要把通知变短，明白办把通知变可执行、可核对。三个设计取向：")
bullet("（a）事实缺就写「通知未说明」，绝不自动补齐；")
bullet("（b）每条关键事实附原文引文，程序校验引文必须真实存在于原文；")
bullet("（c）把「通知明确要求」与「为了问清楚而提的建议」分开标注。")

h("1.4 产品边界（明确不做）", 2)
bullet("不替发布方制定办理要求，不判断资格，不做医疗或法律判断。")
bullet("不代替用户报名、预约或办理任何业务；清单打勾只表示「我自己准备好了」。")
bullet("本版本不含账号系统、历史记录、支付、地图、OCR、语音、多模型切换。")

h("1.5 一句话定位", 2)
para("不是「更懂 AI」，而是更敢让长辈照着做——每一个结论都能点回原文，每一处缺失都如实标出。", bold=True)
doc.add_page_break()

# ================= 二、目标用户 =================
h("二、目标用户", 1, BRAND)
h("2.1 核心用户：60 岁以上社区居民", 2)
para("典型情形：独居或与子女异地居住；会用微信收发消息，但不熟悉在手机上查找、比对信息；对「AI」没有概念，只关心「这条通知我该怎么办」。他们的共同约束不是「不愿学」，而是看不清、点不准、怕弄错。")

h("2.2 次要用户：子女 / 照护者", 2)
para("常替长辈读通知、转发通知。他们需要「能直接转给爸妈」的一段文字，而不是还要再解释一遍的摘要。")

h("2.3 使用约束与设计对策", 2)
table(["使用约束", "设计对策", "当前实现（可核对）"], [
    ["看不清小字", "正文加大、行距放宽、短句", "正文 20px、行高 1.75；关键值 22—24px"],
    ["手指点不准", "触控区放大、留白充足", "按钮最小高度 56px（要求 ≥44px），复选框 30px"],
    ["怕点错、怕弄坏", "关键动作可回退、不自动提交业务", "「换一条通知」随时重置；勾选仅本地状态，不上传、不代办"],
    ["不信 AI", "结论可回到原文核对", "每个事实项「看原文依据」，展开显示原文原句"],
    ["不会用复杂界面", "一屏一件事，只留一个主按钮", "首页只有输入框 +「帮我看懂」+ 三条样例"],
    ["手机屏幕窄", "窄屏优先，不横向滚动", "360px / 390px 视口实测无横向滚动"],
], widths=[3.2, 6.0, 7.0])

h("2.4 关键假设与验证状态", 2)
table(["假设", "状态", "说明"], [
    ["长辈读通知时确实需要「时间/地点/对象/材料」的结构化帮助", "待验证", "来自设计前的需求推理，尚未做用户访谈"],
    ["「可展开原文依据」能提升长辈对结果的信任", "待验证", "尚未做对照测试"],
    ["行动卡复制给子女是有效的求助路径", "待验证", "功能已实现，转发行为未观测"],
    ["20px 字号 + 56px 触控区足以支撑独立完成一次操作", "待验证", "仅完成尺寸层面的技术检查，未做实机与真人测试"],
], widths=[7.6, 2.2, 6.4])
para("以上「待验证」项均为设计假设。本产品尚未开展目标用户效果验证，因此不主张任何使用效果（如「提升理解率」「减少误读」）。", size=9.5, color=GREY)
doc.add_page_break()

# ================= 三、核心功能 =================
h("三、核心功能", 1, BRAND)
table(["编号", "功能", "实现要点"], [
    ["F1", "通知输入", "支持粘贴中文通知，上限 5000 字；空白与超长都有明确提示且不静默截断；内置 3 条明确标注「虚构演示样例」的通知"],
    ["F2", "白话行动卡", "输出事项说明、适用对象、按用途分开的时间、地点、准备事项、下一步行动、待确认问题"],
    ["F3", "原文核对", "每个事实项标注「原文写明 / 通知未说明 / 说得含糊 / 前后不一致」，可展开逐字原文；没有的信息显示「通知未说明」，不自动补齐"],
    ["F4", "准备与转发", "准备事项可勾选；一键复制整张行动卡，复制文本保留缺失标记与「确认建议」标记；勾选仅表示个人准备状态"],
], widths=[1.4, 2.6, 12.2], font=9)

h("3.1 状态不只靠颜色区分", 2)
para("「原文写明」用实线边框 + ✓，「通知未说明」用虚线边框 + —，「说得含糊」用点线边框 + ?，「前后不一致」用双线边框 + !。色弱用户、或在强光下看屏幕，也能分辨四种状态的差别。")

h("3.2 异常与降级：不让长辈卡住，也不给假答案", 2)
table(["情况", "页面行为", "设计理由"], [
    ["超过 5000 字", "明确告知实际字数与上限，提示分两次贴", "不静默截断，避免漏掉关键信息"],
    ["等太久（>45 秒）", "中止请求，说明可能网络慢，给「再试一次」", "不让用户对着转圈猜"],
    ["模型结果对不上原文", "自动重试 1 次；仍不行则显示「没看明白」，不出行动卡", "宁可不给，也不给错的"],
    ["服务不可用 / 配额用尽", "显示具体原因，明确「这不是通知本身的问题」", "不用固定答案冒充 AI"],
    ["剪贴板被禁用", "退到旧接口，再不行弹出手动全选复制框", "复制是核心动作，不能卡死"],
], widths=[3.6, 7.6, 5.0], font=9)

h("3.3 复制出来的文字格式（真实格式）", 2)
for line in [
    "【明白办】行动卡",
    "这条通知说的是：社区要办手机使用指导活动；想参加的人要在 10 月 25 日 17:00 前到一楼服务台报名。",
    "一、谁能参加　仅限已报名居民（原文写明）",
    "二、时间　报名截止：2026 年 10 月 25 日 17:00 前（原文写明）；活动开始：2026 年 10 月 26 日 9:00—11:00（原文写明）",
    "三、地点　社区服务中心二楼活动室（原文写明）",
    "四、要带什么　[ ] 本人手机（原文写明）",
    "五、您接下来做什么　[ ] 在 10 月 25 日 17:00 前到一楼服务台报名（通知明确要求）；[ ] 先确认自己是否已经报上名（确认建议）",
    "—— 提醒 ——　打勾只表示您自己准备好了，不代表已经报名或办好。内容仅作理解辅助，以原文为准。",
]:
    para(line, size=9.5, space_after=1)
doc.add_page_break()

# ================= 四、AI 角色与价值 =================
h("四、AI 角色与价值", 1, BRAND)
para("AI 承担核心的理解与表达工作：读懂任意写法的中文通知（含书面语、简称、模糊表述），判断哪些是事实、哪些缺失、哪些互相打架，再转写成长辈能懂的白话，并保持「原文要求」与「建议」的区分。这无法用固定规则或模板实现——通知的措辞千变万化。")

h("4.1 内置在模型里的十条硬约束", 2)
table(["#", "约束", "为什么必须有"], [
    ["1", "时间、年份、地址、电话、材料、费用、报名方式、资格不得编造", "编造事实比不给答案更危险"],
    ["2", "区分报名截止、活动开始、服务时间；保留「已报名」「仅限」等限制", "这是长辈最容易误读的地方"],
    ["3", "无更新关系的时间冲突必须并列保留并提示确认", "擅自二选一等于替用户做决定"],
    ["4", "无可靠发布日期的「明天/周五」原样保留，不做相对日期换算", "用服务器当天日期换算极易算错"],
    ["5", "原文只写「带手机」就不许添身份证等常见材料", "「合理补充」正是幻觉高发区"],
    ["6", "关键事实必须给逐字原文摘录，引文必须真实存在", "让结论可以被核对"],
    ["7", "原文要求与确认建议分开，不把建议包装成义务", "避免制造不存在的办理要求"],
    ["8", "非通知输入判为 not_notice，不硬凑行动卡", "宁可不给，也不给错的"],
    ["9", "不做医疗建议、法律判断、资格认定，不声称已代办", "越界即风险"],
    ["10", "只返回约定 JSON，不返回推理过程与无关对话", "保证下游可校验、可渲染"],
], widths=[1.0, 9.6, 5.6], font=9)

h("4.2 模型之外的第二道闸：程序校验", 2)
para("模型输出不直接上屏。程序逐项检查：结构是否完整、status 取值是否合法、missing 是否被偷偷填了值、每条引文能否在原文中逐字找到（只忽略空白与全半角标点差异）。不通过则自动重试 1 次；仍不通过，页面显示「没看明白」，不渲染行动卡。")
para("实测佐证：12 条用例中，B2（无发布日期的「明天上午九点」）首次输出未通过校验，系统自动重试后通过——重试机制在真实运行中被触发并生效。", size=9.5, color=GREY)

h("4.3 运行时的模型服务（如实记录）", 2)
table(["项目", "实际情况"], [
    ["接入方式", "WorkBuddy 云服务托管的免密钥大模型通道；浏览器直接调用数据面，服务端按 Origin 校验，前端不保存长期密钥"],
    ["实际模型", "运行时通过模型列表接口读取并选用 deepseek-v4.1-flash（Deepseek-V4.1-Flash），与页面底部「AI 实时生成」的标注一致"],
    ["开发执行模型", "DeepSeek V4.1 Flash（WorkBuddy 开发侧模型，与运行时模型是两条独立链路，不混为一谈）"],
    ["调用参数", "流式输出、temperature 0.1、单次超时 45 秒、失败最多重试 1 次"],
    ["无法确认的部分", "上游提供方的机房与部署形态由平台托管，本项目无法核实，故不作声明"],
], widths=[3.2, 13.0], font=9)
doc.add_page_break()

# ================= 五、技术实现 =================
h("五、技术实现思路", 1, BRAND)
para("整条链路只有一个外部依赖：WorkBuddy 云服务的大模型通道。没有自建后端、没有数据库、没有登录，因此也不存在「服务端存了通知原文」这个问题。浏览器端承担输入校验、请求编排与结果校验；系统提示词保存为工程文件并在运行时读取，保证文档与代码一致。")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(CHART, width=Cm(15.6))
para("图 1　明白办技术实现与数据流（依据实际代码绘制）", size=8.5, color=GREY)

h("5.1 技术选型", 2)
table(["层", "选型", "理由"], [
    ["前端", "原生 HTML + CSS + 原生 JS（无框架、无构建步骤）", "体积小、加载快、无构建链故障点；长辈设备多为中低端手机"],
    ["模型接入", "WorkBuddy 云服务 SDK（浏览器 CDN 全局版）", "免密钥、免自建后端，前端不保存任何长期密钥"],
    ["数据结构", "固定 JSON 契约（10 个顶层字段）", "可被程序逐项校验，避免「一段自然语言没人能验」"],
    ["校验", "自研 js/validate.js（结构与引文回溯）", "把「引文必须真实存在」变成机器可判定的规则"],
    ["托管", "静态站点托管，单端口 HTTP 服务", "无数据库、无登录、无历史记录，攻击面小"],
], widths=[2.4, 6.6, 7.2], font=9)

h("5.2 关键实现细节", 2)
bullet("单一事实来源：系统提示词保存为 app/prompts/system-prompt.txt，页面运行时读取后作为 messages[0] 发送，避免「文档一份、代码一份」的漂移。", size=9.5)
bullet("过期响应丢弃：每次分析递增请求序号，回调先比对序号，旧请求的晚到响应直接丢弃，防止新通知显示旧答案。", size=9.5)
bullet("防重复提交：加载中禁用主按钮，按钮文案切换为「正在看，请稍等…」。", size=9.5)
bullet("纯文本渲染：模型输出与用户输入一律通过 textContent 写入 DOM，通知里的 <script> 或 <img onerror> 不会被执行。", size=9.5)
bullet("剪贴板降级：优先 navigator.clipboard，失败退 execCommand，再失败弹出手动全选复制框。", size=9.5)
bullet("失败不伪装：配置缺失、SDK 未加载、超时、校验不通过，均显示具体原因，绝不用固定答案冒充 AI 结果。", size=9.5)

h("5.3 输出结构（节选）", 2)
para("document_status · summary · audience{value,status,evidence[]} · times[{purpose,value,status,evidence[]}] · location{…} · preparation[…] · actions[{type:\"required|suggested\",value,evidence[]}] · questions[] · warnings[]", size=9.5)
para("status 仅取 stated / missing / unclear / conflict；status 为 missing 时 value 必须为 null、evidence 必须为空数组。完整说明见工程内 app/prompts/data-schema.md。", size=9.5, color=GREY)
doc.add_page_break()

# ================= 六、应用场景 =================
h("六、应用场景示例", 1, BRAND)
para("以下输入为测试用通知（含虚构演示样例），输出为线上版本的实际返回结果摘录，非示意性编写。", size=9.5, color=GREY)

h("6.1 活动报名：把「报名截止」和「活动开始」拆开", 2)
para("输入：请于 2026 年 10 月 25 日 17:00 前报名，活动于 10 月 28 日 9:00 在社区服务中心开始。")
para("输出要点：时间被拆成两条——报名截止「2026 年 10 月 25 日 17:00 前」、活动开始「10 月 28 日 9:00」；地点「社区服务中心」；报名方式原文未写，提示「需要确认」。价值在于：长辈最常见的错误就是把「25 日前报名」当成「25 日去参加」，这里在结构上就分开了。")

h("6.2 服务登记：材料不多不少", 2)
para("输入：社区食堂下周起为 60 周岁以上居民提供助餐服务，每餐 12 元。请携带身份证和户口本，于 10 月 28 日前到社区居委会登记。登记时间为工作日上午 9:00—11:30。")
para("输出要点：对象「60 周岁以上居民」；费用「每餐 12 元」（未说免费）；准备「身份证、户口本」两项，不多不少；登记截止「10 月 28 日前」；登记时段「工作日上午 9:00—11:30」；服务开始「下周起」原样保留（不换算成具体日期）。")

h("6.3 设施维护：没有材料就别让人准备", 2)
para("输入：因电梯年度检修，3 号楼 1 单元电梯将于本周四上午 8 点至下午 5 点停用，检修期间请走楼梯。如有变动，另行通知。")
para("输出要点：时间保留「本周四上午 8 点至下午 5 点」；准备事项为空——原文没要求带任何东西，产品就明确说「通知没提要带什么」，而不是照习惯列一串证件。")

h("6.4 信息不全与不是通知", 2)
para("输入 A：社区将于 10 月 29 日上午 9 点组织志愿者清理小区绿化带，欢迎居民报名参加。")
para("输出 A：地点只还原到「小区绿化带」，并在待确认问题里追问「活动在小区哪个具体位置集合」「通过什么方式报名」；报名截止标为「通知未说明」。")
para("输入 B：今天天气不错，我们下午去打球吧，顺便一起吃个饭。")
para("输出 B：判定为「这看起来不是一条通知」，不输出任何事实，也不生成一张看起来完整的行动卡。")

h("6.5 可延伸的公共服务场景", 2)
para("同一套「事实 + 状态 + 引文」结构可复用到：居委会选举公告、疫苗/体检安排、水电气检修、老年卡年审、补贴申领、活动室开放时间调整。核心不是某一类通知，而是把公文式表述转成可核对的白话行动。")
doc.add_page_break()

# ================= 七、验证与局限 =================
h("七、验证情况、局限与孵化方向", 1, BRAND)
h("7.1 已完成的测试（真实执行，可复现）", 2)
table(["项目", "结果", "说明"], [
    ["12 条通知用例", "12 / 12 通过", "8 条正常通知 + 4 条边界（缺地点、无发布日期的相对时间、无更新关系的时间冲突、非通知文本）。用例与参考答案在运行产品前写好"],
    ["页面行为测试 10 项", "10 / 10 通过", "空白输入、超长不截断、提示注入、连续提交、过期响应、复制降级、HTML 注入渲染、360/390px 布局、触控尺寸、服务不可用不伪造答案"],
    ["单次响应耗时", "1.8—7.4 秒", "12 次调用实测区间，中位数约 4.7 秒；页面超时阈值 45 秒"],
    ["校验重试链路", "触发 1 次并成功", "B2 首次输出未过校验，自动重试后通过"],
    ["窄屏适配", "通过", "360px 与 390px 视口下文档宽度 = 视口宽度，无横向滚动"],
], widths=[3.6, 2.6, 10.0], font=9)
para("测试环境：Windows 11 + Microsoft Edge（无头浏览器自动化），视口 390×844；测试对象为线上真实版本 v1.0.0，走真实模型调用。详细记录见《09_明白办_测试报告.pdf》。通过这 12 条只说明「在当前测试集上符合预期」，不代表任何输入都不会出错。", size=9.5, color=GREY)

h("7.2 尚未做到 / 已知局限", 2)
bullet("未开展目标用户效果验证。尚无真人试用数据，文中不出现任何效果百分比或满意度结论。", size=9.5)
bullet("未做真机测试。当前仅为桌面浏览器模拟手机尺寸与触控，不等于在长辈的真实手机上验证过。", size=9.5)
bullet("相对日期不换算。「明天」「周五」原样保留并提示确认，因为缺少可靠发布日期时换算极易出错。", size=9.5)
bullet("校验是机械检查。能发现引文对不上、结构缺字段，不能保证「引文真实但理解错了」这种情况。", size=9.5)
bullet("仅支持文字。不支持图片、截图、语音输入；长辈收到的多是图片通知，这是当前最大使用门槛。", size=9.5)
bullet("依赖联网与模型服务。网络差或配额用尽时会失败（页面会明确说明，不会给假答案）。", size=9.5)

h("7.3 孵化方向", 2)
bullet("接入社区公众号/网格群的消息入口，减少「复制粘贴」这一步。", size=9.5)
bullet("增加图片通知识别（OCR），覆盖长辈真正的使用场景。", size=9.5)
bullet("为社区工作者提供「写通知时自查」模式：发布前检测时间/地点/材料是否缺失。", size=9.5)
bullet("与社区养老服务站合作，做小规模真人对照测试，验证「原文依据」是否真的提升信任与完成率。", size=9.5)

h("八、AIGC 与素材来源声明", 1, BRAND)
para("本产品的界面代码、系统提示词、数据结构与本文档，由项目组在 AI 编程助手（WorkBuddy，开发执行模型 DeepSeek V4.1 Flash）辅助下完成；产品运行时由大语言模型 DeepSeek-V4.1-Flash 实时生成行动卡。文中截图取自本产品真实运行界面；标注「虚构演示样例」的示例通知为本项目编写，不是任何真实社区的公告。详见《12_明白办_工具素材来源与原创记录.md》。", size=9.5, color=GREY)

doc.save(OUT)
print("saved", OUT)

# 回读校验
d2 = Document(OUT)
n = sum(1 for _ in d2.paragraphs)
chars = sum(len(p.text) for p in d2.paragraphs)
print("段落数", n, "正文字符数", chars, "表格数", len(d2.tables))
